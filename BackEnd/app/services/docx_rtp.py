"""Service untuk menghasilkan Dokumen Rencana Tindak Pengendalian (RTP) dalam format .docx.

Menggunakan:
1. docxtpl untuk memuat dan mengisi template Word utama (Cover, Kata Pengantar,
   Daftar Isi, Bab I s.d. Bab V, dan Cover Kertas Kerja).
2. python-docx untuk menambahkan Lampiran Kertas Kerja (Form 1.a sampai Form 10)
   dengan tabel, styling resmi, orientasi Landscape/Portrait, serta tanda tangan.
"""
from io import BytesIO
import os
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Cm, Pt, RGBColor
from docxtpl import DocxTemplate
from sqlmodel import Session, select

from app.models import (
    Infokom,
    KonteksOperasionalOpd,
    KonteksStrategisOpd,
    MonitoringPi,
    Opd,
    RtpCee,
)
from app.routers import cee, konteks, master, monitoring, risiko
from app.services.excel_sheet import as_lines, numbered

TEMPLATE_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "templates")
TEMPLATE_PATH = os.path.join(TEMPLATE_DIR, "template_rtp.docx")

ALL_FORMS = [
    "f1a", "f1b", "f1c", "f6",
    "f2a", "f2b", "f2c",
    "f3a", "f3b", "f3c",
    "f4", "f5", "f7",
    "f8", "f9", "f10",
]

FORM_LABELS = {
    "f1a": ("Lampiran 5", "Form 1.a", "REKAPITULASI HASIL KUESIONER PENILAIAN LINGKUNGAN PENGENDALIAN INTERN\nCONTROL ENVIRONMENT EVALUATION (CEE)"),
    "f1b": ("Lampiran 5", "Form 1.b", "CEE BERDASARKAN DOKUMEN\nKONDISI KERENTANAN LINGKUNGAN PENGENDALIAN INTERN"),
    "f1c": ("Lampiran 5", "Form 1.c", "SIMPULAN SURVEI PERSEPSI ATAS LINGKUNGAN PENGENDALIAN INTERN"),
    "f6": ("Lampiran 5", "Form 6", "PENILAIAN ATAS KEGIATAN PENGENDALIAN YANG ADA DAN MASIH DIBUTUHKAN\n(RTP ATAS KELEMAHAN LINGKUNGAN PENGENDALIAN / CEE)"),
    "f2a": ("Lampiran 5", "Form 2.a", "PENETAPAN KONTEKS RISIKO STRATEGIS PEMERINTAH DAERAH"),
    "f2b": ("Lampiran 5", "Form 2.b", "PENETAPAN KONTEKS RISIKO STRATEGIS OPD"),
    "f2c": ("Lampiran 6", "Form 2.c", "PENETAPAN KONTEKS RISIKO OPERASIONAL OPD"),
    "f3a": ("Lampiran 5", "Form 3.a", "KERTAS KERJA IDENTIFIKASI RISIKO STRATEGIS PEMERINTAH DAERAH"),
    "f3b": ("Lampiran 5", "Form 3.b", "KERTAS KERJA IDENTIFIKASI RISIKO STRATEGIS OPD"),
    "f3c": ("Lampiran 5", "Form 3.c", "KERTAS KERJA IDENTIFIKASI RISIKO OPERASIONAL OPD"),
    "f4": ("Lampiran 5", "Form 4", "KERTAS KERJA HASIL ANALISIS RISIKO"),
    "f5": ("Lampiran 5", "Form 5", "KERTAS KERJA DAFTAR RISIKO PRIORITAS"),
    "f7": ("Lampiran 5", "Form 7", "PENILAIAN ATAS KEGIATAN PENGENDALIAN YANG ADA DAN MASIH DIBUTUHKAN\n(RTP ATAS HASIL IDENTIFIKASI RISIKO)"),
    "f8": ("Lampiran 5", "Form 8", "RENCANA DAN REALISASI ATAS PENGKOMUNIKASIAN ATAS KEGIATAN PENGENDALIAN YANG DIBANGUN"),
    "f9": ("Lampiran 5", "Form 9", "RENCANA DAN REALISASI PEMANTAUAN ATAS KEGIATAN PENGENDALIAN INTERN YANG DIBUTUHKAN"),
    "f10": ("Lampiran 5", "Form 10", "PENCATATAN KEJADIAN RISIKO (RISK EVENT) DAN PELAKSANAAN RTP"),
}

JENIS_LABEL = {
    "strategis_pemda": "Risiko Strategis Pemda (Form 3.a)",
    "strategis_opd": "Risiko Strategis OPD (Form 3.b)",
    "operasional_opd": "Risiko Operasional OPD (Form 3.c)",
}


def _shade(cell, hex_color="D9E2F3"):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)


def _set_cell_padding(cell, top=70, bottom=70, left=100, right=100):
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    cell._tc.get_or_add_tcPr().append(tcMar)


def _cell_p(cell, text="", bold=False, italic=False, size=9, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = cell.paragraphs[0]
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.05
    if text:
        run = p.add_run(str(text))
        run.bold = bold
        run.italic = italic
        run.font.name = "Times New Roman"
        run.font.size = Pt(size)
    return p


class DocxContext:
    def __init__(self, session: Session, opd_id: int, tahun: int):
        self.s = session
        self.opd_id = opd_id
        self.tahun = tahun
        self.opdctx = master.opd_print_context(opd_id, session) or {}
        self.opd = session.exec(select(Opd).where(Opd.id_sub_pd == opd_id)).first()
        self.opd_name = self.opdctx.get("nama_pd") or (self.opd.nama_pd if self.opd else "-")
        self.nama_kepala = self.opdctx.get("nama_kepala") or (self.opd.nama_kepala if self.opd else "")
        self.nip_kepala = self.opdctx.get("nip_kepala") or (self.opd.nip_kepala if self.opd else "")
        self.jabatan_kepala = self.opdctx.get("jabatan_kepala") or (self.opd.jabatan_kepala if self.opd else f"Kepala {self.opd_name}")
        self.bidang_urusan = self.opdctx.get("bidang_urusan") or "URUSAN PEMERINTAHAN WAJIB YANG BERKAITAN DENGAN PELAYANAN DASAR"


def _add_section(doc, landscape=True):
    sec = doc.add_section(WD_SECTION_START.NEW_PAGE)
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width = Cm(29.7)
        sec.page_height = Cm(21.0)
        sec.top_margin = Cm(1.5)
        sec.bottom_margin = Cm(1.5)
        sec.left_margin = Cm(1.8)
        sec.right_margin = Cm(1.8)
    else:
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        sec.top_margin = Cm(2.0)
        sec.bottom_margin = Cm(2.0)
        sec.left_margin = Cm(2.5)
        sec.right_margin = Cm(2.0)
    return sec


def _add_form_header(doc, lampiran_label: str, form_no: str, title: str, ctx: DocxContext):
    p_lamp = doc.add_paragraph()
    p_lamp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_lamp.paragraph_format.space_after = Pt(2)
    p_lamp.paragraph_format.space_before = Pt(0)
    run_l = p_lamp.add_run(f"{lampiran_label}\n{form_no}")
    run_l.font.name = "Times New Roman"
    run_l.font.size = Pt(10)
    run_l.bold = True

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(2)
    p_title.paragraph_format.space_before = Pt(0)
    run_t = p_title.add_run(title)
    run_t.font.name = "Times New Roman"
    run_t.font.size = Pt(11)
    run_t.bold = True

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(8)
    p_sub.paragraph_format.space_before = Pt(0)
    run_s = p_sub.add_run(f"{ctx.opd_name.upper()} — TAHUN PENILAIAN {ctx.tahun}")
    run_s.font.name = "Times New Roman"
    run_s.font.size = Pt(10)
    run_s.bold = True


def _add_signature(doc, ctx: DocxContext, date_text: str | None = None):
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(8)
    p_space.paragraph_format.space_after = Pt(0)

    # Use a 2-column table for right-aligned signature
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.columns[0].width = Cm(16.0)
    tbl.columns[1].width = Cm(9.0)

    cell_right = tbl.rows[0].cells[1]
    p = cell_right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)

    d_txt = date_text or f"Kota Tegal, Desember {ctx.tahun - 1}"
    run_d = p.add_run(f"{d_txt}\n{ctx.jabatan_kepala}\n\n\n\n")
    run_d.font.name = "Times New Roman"
    run_d.font.size = Pt(10)

    run_n = p.add_run(f"{ctx.nama_kepala or '(………………………………)'}\n")
    run_n.font.name = "Times New Roman"
    run_n.font.size = Pt(10)
    run_n.bold = True

    if ctx.nip_kepala:
        run_nip = p.add_run(f"NIP. {ctx.nip_kepala}")
        run_nip.font.name = "Times New Roman"
        run_nip.font.size = Pt(10)


# ==============================================================================
# FORM BUILDERS
# ==============================================================================

def append_form_1a(doc, ctx: DocxContext):
    _add_section(doc, landscape=True)
    lamp, no, title = FORM_LABELS["f1a"]
    _add_form_header(doc, lamp, no, title, ctx)

    data = cee.form_1a(ctx.opd_id, ctx.tahun, ctx.s)
    responden = data.get("responden") or []
    nresp = len(responden)

    # Headers: No | Pertanyaan / Sub-unsur | R1..Rn | Modus | Simpulan
    n_cols = 2 + (nresp if nresp > 0 else 1) + 2
    tbl = doc.add_table(rows=1, cols=n_cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"

    hdr = tbl.rows[0]
    _cell_p(hdr.cells[0], "No", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _cell_p(hdr.cells[1], "Pertanyaan / Sub-unsur", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _shade(hdr.cells[0])
    _shade(hdr.cells[1])
    _set_cell_padding(hdr.cells[0])
    _set_cell_padding(hdr.cells[1])

    c_idx = 2
    if responden:
        for r in responden:
            kod = r.get("kode_responden") if isinstance(r, dict) else getattr(r, "kode_responden", "")
            _cell_p(hdr.cells[c_idx], kod or "-", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=8)
            _shade(hdr.cells[c_idx])
            _set_cell_padding(hdr.cells[c_idx])
            c_idx += 1
    else:
        _cell_p(hdr.cells[c_idx], "Jawaban", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _shade(hdr.cells[c_idx])
        _set_cell_padding(hdr.cells[c_idx])
        c_idx += 1

    _cell_p(hdr.cells[c_idx], "Modus", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _shade(hdr.cells[c_idx])
    _set_cell_padding(hdr.cells[c_idx])
    c_idx += 1

    _cell_p(hdr.cells[c_idx], "Simpulan", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _shade(hdr.cells[c_idx])
    _set_cell_padding(hdr.cells[c_idx])

    for kat in data.get("kategori") or []:
        # Category header row
        row_kat = tbl.add_row()
        _cell_p(row_kat.cells[0], kat.get("kode"), bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell_p(row_kat.cells[1], kat.get("nama"), bold=True)
        _shade(row_kat.cells[0], "F2F2F2")
        _shade(row_kat.cells[1], "F2F2F2")
        _set_cell_padding(row_kat.cells[0])
        _set_cell_padding(row_kat.cells[1])
        for c in range(2, n_cols - 1):
            _shade(row_kat.cells[c], "F2F2F2")
            _set_cell_padding(row_kat.cells[c])
        _cell_p(row_kat.cells[n_cols - 1], kat.get("simpulan"), bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _shade(row_kat.cells[n_cols - 1], "F2F2F2")
        _set_cell_padding(row_kat.cells[n_cols - 1])

        for i, pq in enumerate(kat.get("pertanyaan") or [], start=1):
            r_data = tbl.add_row()
            _cell_p(r_data.cells[0], str(i), align=WD_ALIGN_PARAGRAPH.CENTER)
            _cell_p(r_data.cells[1], pq.get("pertanyaan"))
            _set_cell_padding(r_data.cells[0])
            _set_cell_padding(r_data.cells[1])

            jawaban = pq.get("jawaban") or {}
            c_pos = 2
            if responden:
                for r in responden:
                    rid = r.get("id") if isinstance(r, dict) else getattr(r, "id", None)
                    v = jawaban.get(rid, jawaban.get(str(rid)))
                    _cell_p(r_data.cells[c_pos], str(v) if v is not None else "-", align=WD_ALIGN_PARAGRAPH.CENTER, size=8)
                    _set_cell_padding(r_data.cells[c_pos])
                    c_pos += 1
            else:
                _cell_p(r_data.cells[c_pos], "-", align=WD_ALIGN_PARAGRAPH.CENTER)
                _set_cell_padding(r_data.cells[c_pos])
                c_pos += 1

            _cell_p(r_data.cells[c_pos], str(pq.get("modus") or "-"), align=WD_ALIGN_PARAGRAPH.CENTER)
            _set_cell_padding(r_data.cells[c_pos])
            c_pos += 1

            _cell_p(r_data.cells[c_pos], str(pq.get("simpulan") or "-"), align=WD_ALIGN_PARAGRAPH.CENTER)
            _set_cell_padding(r_data.cells[c_pos])

    _add_signature(doc, ctx)


def append_form_1b(doc, ctx: DocxContext):
    _add_section(doc, landscape=True)
    lamp, no, title = FORM_LABELS["f1b"]
    _add_form_header(doc, lamp, no, title, ctx)

    rows = cee.form_1b(ctx.opd_id, ctx.tahun, ctx.s) or []
    tbl = doc.add_table(rows=1, cols=5)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"

    hdr = tbl.rows[0]
    cols = ["No", "Sub-unsur / Aspek", "Hasil Reviu", "Sumber Data", "Uraian Kelemahan"]
    for i, h in enumerate(cols):
        _cell_p(hdr.cells[i], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _shade(hdr.cells[i])
        _set_cell_padding(hdr.cells[i])

    for r in rows:
        r_row = tbl.add_row()
        _cell_p(r_row.cells[0], str(r.get("nomor") or ""), align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell_p(r_row.cells[1], r.get("aspek") or "")
        _cell_p(r_row.cells[2], r.get("simpulan") or "", align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell_p(r_row.cells[3], r.get("sumber_data") or "")
        # Kelemahan
        kel_items = [k.get("uraian") for k in (r.get("kelemahan") or []) if k.get("uraian")]
        kel_text = numbered(kel_items) if kel_items else (r.get("keterangan") or "")
        _cell_p(r_row.cells[4], kel_text)
        for c in r_row.cells:
            _set_cell_padding(c)

    _add_signature(doc, ctx)


def append_form_1c(doc, ctx: DocxContext):
    _add_section(doc, landscape=True)
    lamp, no, title = FORM_LABELS["f1c"]
    _add_form_header(doc, lamp, no, title, ctx)

    rows = cee.form_1c(ctx.opd_id, ctx.tahun, ctx.s) or []
    tbl = doc.add_table(rows=1, cols=6)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"

    hdr = tbl.rows[0]
    cols = ["No", "Sub-unsur", "Hasil Persepsi", "Hasil Dokumen", "Simpulan", "Penjelasan"]
    for i, h in enumerate(cols):
        _cell_p(hdr.cells[i], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _shade(hdr.cells[i])
        _set_cell_padding(hdr.cells[i])

    for r in rows:
        r_row = tbl.add_row()
        _cell_p(r_row.cells[0], str(r.get("no") or ""), align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell_p(r_row.cells[1], r.get("sub_unsur") or "")
        _cell_p(r_row.cells[2], r.get("hasil_persepsi") or "", align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell_p(r_row.cells[3], r.get("hasil_dokumen") or "", align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell_p(r_row.cells[4], r.get("simpulan") or "", align=WD_ALIGN_PARAGRAPH.CENTER)
        kel_items = [k.get("uraian") for k in (r.get("kelemahan") or []) if k.get("uraian")]
        kel_text = numbered(kel_items) if kel_items else (r.get("uraian_dokumen") or "")
        _cell_p(r_row.cells[5], kel_text)
        for c in r_row.cells:
            _set_cell_padding(c)

    _add_signature(doc, ctx)


def append_form_6(doc, ctx: DocxContext):
    _add_section(doc, landscape=True)
    lamp, no, title = FORM_LABELS["f6"]
    _add_form_header(doc, lamp, no, title, ctx)

    rows = ctx.s.exec(
        select(RtpCee)
        .where(RtpCee.opd_id == ctx.opd_id, RtpCee.tahun == ctx.tahun)
        .order_by(RtpCee.no_urut)
    ).all()

    tbl = doc.add_table(rows=1, cols=6)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"

    hdr = tbl.rows[0]
    cols = ["No", "Kondisi Kurang Memadai", "Rencana Tindak Pengendalian", "Penanggung Jawab", "Target Waktu", "Realisasi Waktu"]
    for i, h in enumerate(cols):
        _cell_p(hdr.cells[i], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _shade(hdr.cells[i])
        _set_cell_padding(hdr.cells[i])

    groups: dict[str, list] = {}
    for r in rows:
        groups.setdefault(r.aspek_cee or "(Tanpa Sub-unsur)", []).append(r)

    for aspek, items in groups.items():
        grp_row = tbl.add_row()
        _cell_p(grp_row.cells[0], "", align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell_p(grp_row.cells[1], aspek, bold=True)
        _shade(grp_row.cells[1], "F2F2F2")
        for c in range(2, 6):
            _shade(grp_row.cells[c], "F2F2F2")
        for i, r in enumerate(items, start=1):
            r_row = tbl.add_row()
            _cell_p(r_row.cells[0], str(r.no_urut or i), align=WD_ALIGN_PARAGRAPH.CENTER)
            _cell_p(r_row.cells[1], r.kondisi_kerentanan or "")
            _cell_p(r_row.cells[2], r.rencana_tindak or "")
            _cell_p(r_row.cells[3], r.pemilik_penanggung_jawab or "")
            _cell_p(r_row.cells[4], r.target_waktu or "", align=WD_ALIGN_PARAGRAPH.CENTER)
            _cell_p(r_row.cells[5], r.realisasi_waktu or "", align=WD_ALIGN_PARAGRAPH.CENTER)
            for c in r_row.cells:
                _set_cell_padding(c)

    _add_signature(doc, ctx)


def append_form_2a(doc, ctx: DocxContext):
    _add_section(doc, landscape=False)
    lamp, no, title = FORM_LABELS["f2a"]
    _add_form_header(doc, lamp, no, title, ctx)

    obj = konteks.get_pemda(ctx.tahun, ctx.s)
    d = obj.model_dump() if obj else {}

    tbl = doc.add_table(rows=0, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"

    fields = [
        ("Periode yang Dinilai", "periode_dinilai", False),
        ("Visi", "visi", False),
        ("Misi Strategis RPJMD", "misi_strategis", True),
        ("Penetapan Konteks Tujuan Risiko Strategis Pemda", "penetapan_konteks_tujuan", True),
        ("Penetapan Konteks Sasaran Risiko Strategis Pemda", "penetapan_konteks_sasaran", True),
        ("Penetapan Konteks IKU Risiko Strategis Pemda", "penetapan_konteks_iku", True),
        ("Prioritas Pembangunan Daerah", "prioritas_pembangunan", True),
        ("Program Prioritas", "prioritas_program", True),
    ]

    for label, key, multi in fields:
        raw = d.get(key)
        val = numbered([x for x in as_lines(raw).split("\n") if x]) if multi else (raw or "-")
        r_row = tbl.add_row()
        _cell_p(r_row.cells[0], label, bold=True)
        _cell_p(r_row.cells[1], ":", align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell_p(r_row.cells[2], val)
        _set_cell_padding(r_row.cells[0])
        _set_cell_padding(r_row.cells[1])
        _set_cell_padding(r_row.cells[2])

    tbl.columns[0].width = Cm(6.5)
    tbl.columns[1].width = Cm(0.8)
    tbl.columns[2].width = Cm(11.0)

    _add_signature(doc, ctx)


def append_form_2b(doc, ctx: DocxContext):
    _add_section(doc, landscape=False)
    lamp, no, title = FORM_LABELS["f2b"]
    _add_form_header(doc, lamp, no, title, ctx)

    rows = ctx.s.exec(
        select(KonteksStrategisOpd).where(
            KonteksStrategisOpd.opd_id == ctx.opd_id,
            KonteksStrategisOpd.tahun == ctx.tahun,
        ).order_by(KonteksStrategisOpd.id)
    ).all()
    tujuan = master.list_renstra_tujuan(ctx.opd_id, ctx.tahun, ctx.s)
    sasaran = master.list_renstra_sasaran(ctx.opd_id, ctx.tahun, ctx.s)
    iku = master.list_renstra_iku(ctx.opd_id, ctx.s)
    program = master.list_renstra_program(ctx.opd_id, ctx.s)

    def uniq(vals):
        out = []
        for v in vals:
            if v and v not in out:
                out.append(v)
        return out

    tbl = doc.add_table(rows=0, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"

    items_green = [
        ("Nama Pemda", "Pemerintah Kota Tegal"),
        ("Tahun Penilaian", str(ctx.tahun)),
        ("Periode yang Dinilai", "; ".join(uniq([r.periode_dinilai for r in rows])) or f"Tahun {ctx.tahun}"),
        ("Bidang Urusan", ctx.bidang_urusan),
        ("OPD yang Dinilai", ctx.opd_name),
        ("Sumber Data", "; ".join(uniq([r.sumber_data for r in rows])) or "Renstra OPD"),
        ("Tujuan Strategis", "\n".join(o["value"] for o in tujuan) or "-"),
        ("Sasaran Strategis", "\n".join(o["value"] for o in sasaran) or "-"),
        ("IKU Renstra OPD", numbered([o["value"] for o in iku]) or "-"),
        ("Program", numbered([o["value"] for o in program]) or "-"),
    ]

    for lbl, val in items_green:
        r = tbl.add_row()
        _cell_p(r.cells[0], lbl, bold=True)
        _cell_p(r.cells[1], ":", align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell_p(r.cells[2], val)
        _shade(r.cells[0], "E2EFDA")  # Green tint for renstra
        _shade(r.cells[1], "E2EFDA")
        _shade(r.cells[2], "E2EFDA")
        _set_cell_padding(r.cells[0])
        _set_cell_padding(r.cells[1])
        _set_cell_padding(r.cells[2])

    en_t = uniq([r.tujuan_strategis for r in rows])
    en_s = uniq([r.sasaran_strategis for r in rows])
    en_i = uniq([r.iku_renstra for r in rows])
    en_p = uniq([r.program for r in rows])
    blok_kuning = "\n".join([
        "Tujuan Strategis :", "\n".join(en_t) or "-",
        "Sasaran Strategis :", "\n".join(f"- {s}" for s in en_s) or "-",
        "IKU Renstra OPD :", numbered(en_i) or "-",
        "Program :", numbered(en_p) or "-",
    ])

    r_kuning = tbl.add_row()
    _cell_p(r_kuning.cells[0], "Tujuan, Sasaran, IKU dan Program yang akan dilakukan penilaian risiko", bold=True)
    _cell_p(r_kuning.cells[1], ":", align=WD_ALIGN_PARAGRAPH.CENTER)
    _cell_p(r_kuning.cells[2], blok_kuning)
    _shade(r_kuning.cells[0], "FFF2CC")  # Yellow tint
    _shade(r_kuning.cells[1], "FFF2CC")
    _shade(r_kuning.cells[2], "FFF2CC")
    _set_cell_padding(r_kuning.cells[0])
    _set_cell_padding(r_kuning.cells[1])
    _set_cell_padding(r_kuning.cells[2])

    tbl.columns[0].width = Cm(6.5)
    tbl.columns[1].width = Cm(0.8)
    tbl.columns[2].width = Cm(11.0)

    _add_signature(doc, ctx)


def append_form_2c(doc, ctx: DocxContext):
    _add_section(doc, landscape=True)
    lamp, no, title = FORM_LABELS["f2c"]
    _add_form_header(doc, lamp, no, title, ctx)

    renja = master.list_renja_subkegiatan(ctx.opd_id, ctx.tahun, ctx.s) or {}
    rows = ctx.s.exec(
        select(KonteksOperasionalOpd).where(
            KonteksOperasionalOpd.opd_id == ctx.opd_id,
            KonteksOperasionalOpd.tahun == ctx.tahun,
        )
    ).all()
    checked_sub = {str(r.ref_subkegiatan) for r in rows if r.ref_subkegiatan}
    checked_ind = {str(r.ref_indikator) for r in rows if r.ref_indikator}

    tbl = doc.add_table(rows=1, cols=5)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"

    hdr = tbl.rows[0]
    cols = ["No", "Program / Kegiatan / Sub Kegiatan", "Indikator Program / Kegiatan / Sub Kegiatan", "Target", "Dilakukan Penilaian Risiko"]
    for i, h in enumerate(cols):
        _cell_p(hdr.cells[i], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _shade(hdr.cells[i])
        _set_cell_padding(hdr.cells[i])

    prog: dict[str, dict] = {}
    def ensure_prog(pid, nama):
        return prog.setdefault(str(pid), {"nm": nama, "ind": [], "keg": {}})
    def ensure_keg(g, kid, nama):
        return g["keg"].setdefault(str(kid), {"nm": nama, "ind": [], "subs": []})

    for s in renja.get("subkegiatan") or []:
        ensure_keg(ensure_prog(s["idprogram"], s["nm_program"]), s["idkegiatan"], s["nm_kegiatan"])["subs"].append(s)
    for pi in renja.get("indikator_program") or []:
        ensure_prog(pi["idprogram"], pi["nm_program"])["ind"].append(pi)
    for ki in renja.get("indikator_kegiatan") or []:
        ensure_keg(ensure_prog(ki["idprogram"], ki["nm_program"]), ki["idkegiatan"], ki["nm_kegiatan"])["ind"].append(ki)

    num = 0
    def emit(nama, indikator, target, sel, level):
        nonlocal num
        num += 1
        r = tbl.add_row()
        _cell_p(r.cells[0], str(num), align=WD_ALIGN_PARAGRAPH.CENTER)
        indent = "    " * (level - 1)
        _cell_p(r.cells[1], f"{indent}{nama}" if nama else "", bold=level < 3)
        _cell_p(r.cells[2], indikator or "")
        _cell_p(r.cells[3], target or "", align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell_p(r.cells[4], "✓" if sel else "", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
        for c in r.cells:
            _set_cell_padding(c)

    for g in prog.values():
        if g["ind"]:
            for i, pi in enumerate(g["ind"]):
                emit(g["nm"] if i == 0 else "", pi.get("tolok"), pi.get("target"), str(pi.get("id")) in checked_ind, 1)
        else:
            emit(g["nm"], "", "", False, 1)
        for k in g["keg"].values():
            if k["ind"]:
                for i, ki in enumerate(k["ind"]):
                    emit(k["nm"] if i == 0 else "", ki.get("tolok"), ki.get("target"), str(ki.get("id")) in checked_ind, 2)
            else:
                emit(k["nm"], "", "", False, 2)
            for s in k["subs"]:
                inds = s.get("indikator_list") or [{"tolok": "", "target": ""}]
                nama = f"{s.get('kode_sub_kegiatan') or ''} {s.get('nm_sub_kegiatan') or ''}".strip()
                for i, ind in enumerate(inds):
                    emit(nama if i == 0 else "", ind.get("tolok"), ind.get("target"), i == 0 and str(s["idsubkegiatan"]) in checked_sub, 3)

    _add_signature(doc, ctx)


def _append_form_3(doc, ctx: DocxContext, jenis: str, key: str):
    _add_section(doc, landscape=True)
    lamp, no, title = FORM_LABELS[key]
    _add_form_header(doc, lamp, no, title, ctx)

    groups = risiko.form_4(ctx.opd_id, ctx.tahun, ctx.s)
    items = next((g["items"] for g in groups if g["jenis"] == jenis), [])
    operasional = jenis == "operasional_opd"

    cols = ["No", "Tujuan / Sasaran / Kegiatan", "Indikator Kinerja / Keluaran"]
    if operasional:
        cols.append("Tahap")
    cols += ["Uraian Risiko", "Kode Risiko", "Pemilik Risiko", "Sebab (Uraian)", "Sebab (Sumber)", "C/UC", "Dampak (Uraian)", "Pihak Terkena"]

    tbl = doc.add_table(rows=1, cols=len(cols))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"

    hdr = tbl.rows[0]
    for i, h in enumerate(cols):
        _cell_p(hdr.cells[i], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _shade(hdr.cells[i])
        _set_cell_padding(hdr.cells[i])

    for i, it in enumerate(items, start=1):
        r = tbl.add_row()
        c_i = 0
        _cell_p(r.cells[c_i], str(it.get("no_urut") or i), align=WD_ALIGN_PARAGRAPH.CENTER)
        c_i += 1
        _cell_p(r.cells[c_i], it.get("kegiatan") if operasional else it.get("tujuan_sasaran") or "")
        c_i += 1
        _cell_p(r.cells[c_i], it.get("indikator_keluaran") if operasional else it.get("indikator_kinerja") or "")
        c_i += 1
        if operasional:
            _cell_p(r.cells[c_i], it.get("tahap_kegiatan") or "")
            c_i += 1
        uraian = it.get("uraian_risiko") or ""
        if it.get("keterangan_risiko"):
            uraian = f"{uraian}\n({it.get('keterangan_risiko')})"
        _cell_p(r.cells[c_i], uraian)
        c_i += 1
        _cell_p(r.cells[c_i], it.get("kode_risiko") or "", align=WD_ALIGN_PARAGRAPH.CENTER)
        c_i += 1
        _cell_p(r.cells[c_i], it.get("pemilik_risiko") or "")
        c_i += 1
        _cell_p(r.cells[c_i], it.get("sebab_uraian") or "")
        c_i += 1
        _cell_p(r.cells[c_i], it.get("sebab_sumber") or "", align=WD_ALIGN_PARAGRAPH.CENTER)
        c_i += 1
        _cell_p(r.cells[c_i], it.get("cuc") or "", align=WD_ALIGN_PARAGRAPH.CENTER)
        c_i += 1
        _cell_p(r.cells[c_i], it.get("dampak_uraian") or "")
        c_i += 1
        _cell_p(r.cells[c_i], it.get("dampak_pihak_terkena") or "")

        for cell in r.cells:
            _set_cell_padding(cell)

    _add_signature(doc, ctx)


def append_form_3a(doc, ctx: DocxContext):
    _append_form_3(doc, ctx, "strategis_pemda", "f3a")


def append_form_3b(doc, ctx: DocxContext):
    _append_form_3(doc, ctx, "strategis_opd", "f3b")


def append_form_3c(doc, ctx: DocxContext):
    _append_form_3(doc, ctx, "operasional_opd", "f3c")


def append_form_4(doc, ctx: DocxContext):
    _add_section(doc, landscape=True)
    lamp, no, title = FORM_LABELS["f4"]
    _add_form_header(doc, lamp, no, title, ctx)

    groups = risiko.form_4(ctx.opd_id, ctx.tahun, ctx.s)
    tbl = doc.add_table(rows=1, cols=7)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"

    hdr = tbl.rows[0]
    cols = ["No", "Risiko", "Kode Risiko", "Dampak", "Kemungkinan", "Skala", "Level"]
    for i, h in enumerate(cols):
        _cell_p(hdr.cells[i], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _shade(hdr.cells[i])
        _set_cell_padding(hdr.cells[i])

    for g in groups:
        grp_row = tbl.add_row()
        _cell_p(grp_row.cells[0], "", align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell_p(grp_row.cells[1], JENIS_LABEL.get(g["jenis"], g["jenis"]), bold=True)
        _shade(grp_row.cells[1], "F2F2F2")
        for c in range(2, 7):
            _shade(grp_row.cells[c], "F2F2F2")

        for i, it in enumerate(g.get("items") or [], start=1):
            a = it.get("analisis") or {}
            r = tbl.add_row()
            _cell_p(r.cells[0], str(i), align=WD_ALIGN_PARAGRAPH.CENTER)
            _cell_p(r.cells[1], it.get("uraian_risiko") or "")
            _cell_p(r.cells[2], it.get("kode_risiko") or "", align=WD_ALIGN_PARAGRAPH.CENTER)
            _cell_p(r.cells[3], str(a.get("skala_dampak") or "-"), align=WD_ALIGN_PARAGRAPH.CENTER)
            _cell_p(r.cells[4], str(a.get("skala_kemungkinan") or "-"), align=WD_ALIGN_PARAGRAPH.CENTER)
            _cell_p(r.cells[5], str(a.get("skala_risiko") or "-"), align=WD_ALIGN_PARAGRAPH.CENTER)
            _cell_p(r.cells[6], str(a.get("level") or "-"), align=WD_ALIGN_PARAGRAPH.CENTER)
            for c in r.cells:
                _set_cell_padding(c)

    _add_signature(doc, ctx)


def append_form_5(doc, ctx: DocxContext):
    _add_section(doc, landscape=True)
    lamp, no, title = FORM_LABELS["f5"]
    _add_form_header(doc, lamp, no, title, ctx)

    groups = risiko.form_5(ctx.opd_id, ctx.tahun, ctx.s)
    tbl = doc.add_table(rows=1, cols=7)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"

    hdr = tbl.rows[0]
    cols = ["No", "Risiko Prioritas", "Kode Risiko", "Skala", "Pemilik Risiko", "Penyebab", "Dampak"]
    for i, h in enumerate(cols):
        _cell_p(hdr.cells[i], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _shade(hdr.cells[i])
        _set_cell_padding(hdr.cells[i])

    for g in groups:
        if not g.get("items"):
            continue
        grp_row = tbl.add_row()
        _cell_p(grp_row.cells[0], "", align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell_p(grp_row.cells[1], JENIS_LABEL.get(g["jenis"], g["jenis"]), bold=True)
        _shade(grp_row.cells[1], "F2F2F2")
        for c in range(2, 7):
            _shade(grp_row.cells[c], "F2F2F2")

        for i, it in enumerate(g.get("items") or [], start=1):
            a = it.get("analisis") or {}
            r = tbl.add_row()
            _cell_p(r.cells[0], str(i), align=WD_ALIGN_PARAGRAPH.CENTER)
            _cell_p(r.cells[1], it.get("uraian_risiko") or "")
            _cell_p(r.cells[2], it.get("kode_risiko") or "", align=WD_ALIGN_PARAGRAPH.CENTER)
            _cell_p(r.cells[3], str(a.get("skala_risiko") or "-"), align=WD_ALIGN_PARAGRAPH.CENTER)
            _cell_p(r.cells[4], it.get("pemilik_risiko") or "")
            _cell_p(r.cells[5], it.get("sebab_uraian") or "")
            _cell_p(r.cells[6], it.get("dampak_uraian") or "")
            for c in r.cells:
                _set_cell_padding(c)

    _add_signature(doc, ctx)


def append_form_7(doc, ctx: DocxContext):
    _add_section(doc, landscape=True)
    lamp, no, title = FORM_LABELS["f7"]
    _add_form_header(doc, lamp, no, title, ctx)

    rows = risiko.list_rtp(ctx.opd_id, ctx.tahun, ctx.s) or []
    tbl = doc.add_table(rows=1, cols=8)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"

    hdr = tbl.rows[0]
    cols = ["No", "Risiko Prioritas", "Kode Risiko", "Pemilik", "Pengendalian yang Ada", "Celah Pengendalian", "Rencana Tindak Pengendalian", "Target Waktu"]
    for i, h in enumerate(cols):
        _cell_p(hdr.cells[i], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _shade(hdr.cells[i])
        _set_cell_padding(hdr.cells[i])

    for i, it in enumerate(rows, start=1):
        rtp = it.get("rtp") or {}
        r = tbl.add_row()
        _cell_p(r.cells[0], str(i), align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell_p(r.cells[1], it.get("uraian_risiko") or "")
        _cell_p(r.cells[2], it.get("kode_risiko") or "", align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell_p(r.cells[3], it.get("pemilik_risiko") or "")
        _cell_p(r.cells[4], rtp.get("pengendalian_ada") or "")
        _cell_p(r.cells[5], rtp.get("celah_pengendalian") or "")
        _cell_p(r.cells[6], rtp.get("rencana_tindak") or "")
        _cell_p(r.cells[7], rtp.get("target_waktu") or "", align=WD_ALIGN_PARAGRAPH.CENTER)
        for c in r.cells:
            _set_cell_padding(c)

    _add_signature(doc, ctx)


def append_form_8(doc, ctx: DocxContext):
    _add_section(doc, landscape=True)
    lamp, no, title = FORM_LABELS["f8"]
    _add_form_header(doc, lamp, no, title, ctx)

    rows = ctx.s.exec(
        select(Infokom)
        .where(Infokom.opd_id == ctx.opd_id, Infokom.tahun == ctx.tahun)
        .order_by(Infokom.no_urut)
    ).all()

    tbl = doc.add_table(rows=1, cols=7)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"

    hdr = tbl.rows[0]
    cols = ["No", "Kegiatan Pengendalian", "Media / Bentuk", "Penyedia Informasi", "Penerima Informasi", "Rencana Waktu", "Realisasi Waktu"]
    for i, h in enumerate(cols):
        _cell_p(hdr.cells[i], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _shade(hdr.cells[i])
        _set_cell_padding(hdr.cells[i])

    for i, r in enumerate(rows, start=1):
        row = tbl.add_row()
        _cell_p(row.cells[0], str(r.no_urut or i), align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell_p(row.cells[1], r.kegiatan_pengendalian or "")
        _cell_p(row.cells[2], r.media_bentuk or "")
        _cell_p(row.cells[3], r.penyedia_informasi or "")
        _cell_p(row.cells[4], r.penerima_informasi or "")
        _cell_p(row.cells[5], r.rencana_waktu or "", align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell_p(row.cells[6], r.realisasi_waktu or "", align=WD_ALIGN_PARAGRAPH.CENTER)
        for c in row.cells:
            _set_cell_padding(c)

    _add_signature(doc, ctx)


def _group_by_jenis(rows, key="jenis_risiko"):
    buckets: dict[str, list] = {}
    for r in rows:
        j = r.get(key) if isinstance(r, dict) else getattr(r, key, None)
        buckets.setdefault(j if j in JENIS_LABEL else "_lain", []).append(r)
    out = [(JENIS_LABEL[j], buckets[j]) for j in list(JENIS_LABEL) if buckets.get(j)]
    if buckets.get("_lain"):
        out.append(("Lainnya", buckets["_lain"]))
    return out


def append_form_9(doc, ctx: DocxContext):
    _add_section(doc, landscape=True)
    lamp, no, title = FORM_LABELS["f9"]
    _add_form_header(doc, lamp, no, title, ctx)

    rows = ctx.s.exec(
        select(MonitoringPi)
        .where(MonitoringPi.opd_id == ctx.opd_id, MonitoringPi.tahun == ctx.tahun)
        .order_by(MonitoringPi.no_urut)
    ).all()

    tbl = doc.add_table(rows=1, cols=6)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"

    hdr = tbl.rows[0]
    cols = ["No", "Kegiatan Pengendalian", "Metode Pemantauan", "Penanggung Jawab", "Rencana Waktu", "Realisasi Waktu"]
    for i, h in enumerate(cols):
        _cell_p(hdr.cells[i], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _shade(hdr.cells[i])
        _set_cell_padding(hdr.cells[i])

    groups = _group_by_jenis(rows)
    for label, items in groups:
        grp_row = tbl.add_row()
        _cell_p(grp_row.cells[0], "", align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell_p(grp_row.cells[1], label, bold=True)
        _shade(grp_row.cells[1], "F2F2F2")
        for c in range(2, 6):
            _shade(grp_row.cells[c], "F2F2F2")

        for i, r in enumerate(items, start=1):
            row = tbl.add_row()
            _cell_p(row.cells[0], str(r.no_urut or i), align=WD_ALIGN_PARAGRAPH.CENTER)
            _cell_p(row.cells[1], r.kegiatan_pengendalian or "")
            _cell_p(row.cells[2], r.metode_pemantauan or "")
            _cell_p(row.cells[3], r.penanggung_jawab or "")
            _cell_p(row.cells[4], r.rencana_waktu or "", align=WD_ALIGN_PARAGRAPH.CENTER)
            _cell_p(row.cells[5], r.realisasi_waktu or "", align=WD_ALIGN_PARAGRAPH.CENTER)
            for c in row.cells:
                _set_cell_padding(c)

    _add_signature(doc, ctx)


def append_form_10(doc, ctx: DocxContext):
    _add_section(doc, landscape=True)
    lamp, no, title = FORM_LABELS["f10"]
    _add_form_header(doc, lamp, no, title, ctx)

    raw_rows = [r for r in monitoring.list_event(ctx.opd_id, ctx.tahun, ctx.s) if r.get("events")]
    tbl = doc.add_table(rows=1, cols=7)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"

    hdr = tbl.rows[0]
    cols = ["Kode Risiko", "Risiko", "Tanggal Terjadi", "Sebab Kejadian", "Dampak Kejadian", "RTP", "Realisasi RTP"]
    for i, h in enumerate(cols):
        _cell_p(hdr.cells[i], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _shade(hdr.cells[i])
        _set_cell_padding(hdr.cells[i])

    groups = _group_by_jenis(raw_rows)
    for label, items in groups:
        grp_row = tbl.add_row()
        _cell_p(grp_row.cells[0], label, bold=True)
        _shade(grp_row.cells[0], "F2F2F2")
        for c in range(1, 7):
            _shade(grp_row.cells[c], "F2F2F2")

        for r in items:
            for ev in r.get("events") or []:
                row = tbl.add_row()
                _cell_p(row.cells[0], r.get("kode_risiko") or "", align=WD_ALIGN_PARAGRAPH.CENTER)
                _cell_p(row.cells[1], r.get("uraian_risiko") or "")
                _cell_p(row.cells[2], ev.get("tanggal_terjadi") or "", align=WD_ALIGN_PARAGRAPH.CENTER)
                _cell_p(row.cells[3], ev.get("sebab_kejadian") or "")
                _cell_p(row.cells[4], ev.get("dampak_kejadian") or "")
                _cell_p(row.cells[5], ev.get("rtp") or "")
                _cell_p(row.cells[6], ev.get("realisasi_pelaksanaan_rtp") or "")
                for c in row.cells:
                    _set_cell_padding(c)

    _add_signature(doc, ctx, f"Kota Tegal, {ctx.tahun}")


FORM_APPENDERS = {
    "f1a": append_form_1a,
    "f1b": append_form_1b,
    "f1c": append_form_1c,
    "f6": append_form_6,
    "f2a": append_form_2a,
    "f2b": append_form_2b,
    "f2c": append_form_2c,
    "f3a": append_form_3a,
    "f3b": append_form_3b,
    "f3c": append_form_3c,
    "f4": append_form_4,
    "f5": append_form_5,
    "f7": append_form_7,
    "f8": append_form_8,
    "f9": append_form_9,
    "f10": append_form_10,
}


def generate_rtp_docx(
    session: Session,
    opd_id: int,
    tahun: int,
    forms: list[str] | None = None,
    include_main_doc: bool = True,
) -> BytesIO:
    """Generate complete RTP docx file using docxtpl and python-docx."""
    ctx = DocxContext(session, opd_id, tahun)

    # Context dictionary for Jinja2 template placeholders
    jinja_context = {
        "nama_opd": ctx.opd_name,
        "nama_opd_upper": ctx.opd_name.upper(),
        "tahun": str(tahun),
        "tahun_sebelumnya": str(tahun - 1),
        "nama_pemda": "PEMERINTAH KOTA TEGAL",
        "nama_kepala": ctx.nama_kepala or "(………………………………)",
        "nip_kepala": ctx.nip_kepala or "-",
        "jabatan_kepala": ctx.jabatan_kepala or f"Kepala {ctx.opd_name}",
        "bulan_ttd": "Oktober",
        "bidang_urusan": ctx.bidang_urusan,
        "periode_dinilai": f"Periode Renstra Tahun {tahun-1}-{tahun+2}",
    }

    # Stage 1: Render the base Word template with docxtpl
    if include_main_doc and os.path.exists(TEMPLATE_PATH):
        tpl = DocxTemplate(TEMPLATE_PATH)
        tpl.render(jinja_context)
        buf1 = BytesIO()
        tpl.save(buf1)
        buf1.seek(0)
        doc = Document(buf1)
    else:
        doc = Document()
        doc.add_heading(f"DOKUMEN RENCANA TINDAK PENGENDALIAN (RTP) - {ctx.opd_name}", 0)

    # Stage 2: Append requested Lampiran Forms
    wanted_forms = [k for k in ALL_FORMS if not forms or k in forms]
    for k in wanted_forms:
        appender = FORM_APPENDERS.get(k)
        if appender:
            appender(doc, ctx)

    # Save final document
    out_buf = BytesIO()
    doc.save(out_buf)
    out_buf.seek(0)
    return out_buf
