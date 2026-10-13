"""Template and bounded XLSX parsing for prospective respondents."""
from io import BytesIO
from zipfile import ZipFile

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

KATEGORI_JABATAN = ['Pejabat Struktural', 'Pejabat Fungsional Tertentu', 'Pejabat Fungsional Umum / Staf']

HEADERS = ['NIP', 'Nama', 'Jabatan', 'Email', 'No HP', 'Kategori Jabatan']
FIELDS = ['nip', 'nama_responden', 'jabatan', 'email', 'no_hp', 'kategori_jabatan']
MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 1000


def template():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Responden'
    ws.append(HEADERS)
    ws.freeze_panes = 'A2'
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='334155')
    for col, width in zip('ABCDEF', (25, 35, 35, 40, 22, 40)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2, max_row=MAX_ROWS + 1, max_col=6):
        for cell in row:
            cell.number_format = '@'
    from openpyxl.worksheet.datavalidation import DataValidation
    validation = DataValidation(type='list', formula1='"' + ','.join(KATEGORI_JABATAN) + '"', allow_blank=True)
    validation.errorTitle = 'Kategori tidak valid'
    validation.error = 'Pilih kategori jabatan dari daftar.'
    validation.showErrorMessage = True
    ws.add_data_validation(validation)
    validation.add(f'F2:F{MAX_ROWS + 1}')
    notes = wb.create_sheet('Petunjuk')
    for line in (
        'Isi sheet Responden; jangan mengubah judul kolom pada baris pertama.',
        'Kategori Jabatan boleh kosong untuk data lama; kolom lainnya wajib. Maksimal 1.000 responden per file (.xlsx, 5 MB).',
        'NIP dan No HP harus berupa teks agar digit panjang dan angka nol awal tidak hilang.',
        'OPD dan tahun mengikuti pilihan di aplikasi. Token dibuat otomatis.',
        'Baris valid disimpan; baris tidak valid atau duplikat dilaporkan dan dilewati.',
        'Impor tidak mengubah data lama dan tidak mengirim email undangan.',
    ):
        notes.append([line])
    notes.column_dimensions['A'].width = 110
    output = BytesIO()
    wb.save(output)
    wb.close()
    return output.getvalue()


def read_rows(content):
    if not content or len(content) > MAX_BYTES:
        raise HTTPException(400, 'File Excel kosong atau melebihi 5 MB')
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 20 * 1024 * 1024:
                raise HTTPException(400, 'Isi file Excel terlalu besar')
        wb = load_workbook(BytesIO(content), read_only=True, data_only=False)
        try:
            ws = wb['Responden'] if 'Responden' in wb.sheetnames else wb.worksheets[0]
            rows = ws.iter_rows(max_row=MAX_ROWS + 2, max_col=6)
            header = next(rows, ())
            labels = [str(c.value or '').strip().lower() for c in header]
            if labels not in ([h.lower() for h in HEADERS], [h.lower() for h in HEADERS[:5]] + ['']):
                raise HTTPException(400, 'Kolom harus berurutan: NIP, Nama, Jabatan, Email, No HP, Kategori Jabatan. Gunakan template.')
            result = []
            for number, cells in enumerate(rows, 2):
                if not any(c.value is not None and str(c.value).strip() for c in cells):
                    continue
                if number > MAX_ROWS + 1:
                    raise HTTPException(400, 'Maksimal 1.000 baris data; pisahkan file Excel')
                error = None
                if any(c.data_type == 'f' for c in cells):
                    error = 'Rumus Excel tidak didukung; gunakan nilai teks'
                elif any(c.value is not None and not isinstance(c.value, str) for c in (cells[0], cells[4])):
                    error = 'NIP dan No HP harus berupa teks; isi ulang dari data asli agar digit tetap utuh'
                result.append((number, dict(zip(FIELDS, (c.value for c in cells))), error))
            if ws.max_row and ws.max_row > MAX_ROWS + 1:
                # Reject data beyond the limit even if earlier rows are blank.
                for cells in ws.iter_rows(min_row=MAX_ROWS + 2, max_col=6):
                    if any(c.value is not None for c in cells):
                        raise HTTPException(400, 'Maksimal 1.000 baris data; pisahkan file Excel')
            return result
        finally:
            wb.close()
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(400, 'File tidak dapat dibaca. Gunakan file Excel .xlsx yang valid') from exc
