from io import BytesIO
from openpyxl import Workbook, load_workbook
from fastapi import HTTPException
from sqlmodel import select
import unittest
import test_cee_tokens
from app.models import CeeResponden
from app.routers.cee import import_responden
from app.services.cee_excel import template, read_rows, HEADERS


def excel(rows):
    wb = Workbook()
    wb.active.append(HEADERS)
    for row in rows:
        wb.active.append(row)
    stream = BytesIO()
    wb.save(stream)
    wb.close()
    return stream.getvalue()


class ExcelTests(unittest.TestCase):
    setUp = test_cee_tokens.TokenSurveyTests.setUp
    tearDown = test_cee_tokens.TokenSurveyTests.tearDown

    def test_template(self):
        data = template()
        wb = load_workbook(BytesIO(data))
        self.assertEqual(wb['Responden']['A2'].number_format, '@')
        self.assertEqual(wb['Responden']['E2'].number_format, '@')
        wb.close()
        self.assertEqual(read_rows(data), [])

    def test_import_partial_and_repeated(self):
        data = excel([
            ['001234567890123456', 'Nama', 'Staf', 'new@example.com', '08123456789'],
            ['001234567890123456', 'Duplikat', 'Staf', 'other@example.com', '08123456789'],
            ['002', 'Email salah', 'Staf', 'invalid', '08123456789'],
            [123456789012345678, 'Angka', 'Staf', 'num@example.com', '08123456789'],
            ['003', 'Formula', 'Staf', '=A1', '08123456789'],
            [None, None, None, None, None],
        ])
        result = import_responden(1, 2027, data, self.session, {'role_id': 1})
        self.assertEqual(result['berhasil'], 1)
        self.assertEqual([e['baris'] for e in result['gagal']], [3, 4, 5, 6])
        row = self.session.exec(select(CeeResponden).where(CeeResponden.email == 'new@example.com')).one()
        self.assertEqual(row.nip, '001234567890123456')
        self.assertTrue(row.token_id)
        self.assertEqual(row.kode_responden, 'R2')
        self.assertEqual(import_responden(1, 2027, data, self.session, {'role_id': 1})['berhasil'], 0)

    def test_job_category(self):
        data = excel([
            ['004', 'Kategori Valid', 'Staf', 'valid@example.com', '08123456789', 'Pejabat Struktural'],
            ['005', 'Kategori Salah', 'Staf', 'wrong@example.com', '08123456789', 'Lainnya'],
        ])
        result = import_responden(1, 2027, data, self.session, {'role_id': 1})
        self.assertEqual(result['berhasil'], 1)
        self.assertEqual(result['gagal'][0]['baris'], 3)
        row = self.session.exec(select(CeeResponden).where(CeeResponden.email == 'valid@example.com')).one()
        self.assertEqual(row.kategori_jabatan, 'Pejabat Struktural')

    def test_invalid_files(self):
        for data in (b'not excel', b'', b'x' * (5 * 1024 * 1024 + 1)):
            with self.assertRaises(HTTPException):
                read_rows(data)
        with self.assertRaises(HTTPException):
            read_rows(excel([['1', 'Name', 'Staff', 'a@example.com', '0812345678']] * 1001))
